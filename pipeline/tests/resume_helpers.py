"""Synthetic resume facts only; no real user content in test fixtures."""

import json
from pathlib import Path
from typing import Any
from uuid import uuid4

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from postgres_helpers import PostgresStore

from pipeline.llm.client import Reply
from pipeline.resume.layout import create_template
from pipeline.resume.schema import MasterResume
from pipeline.resume.store import BUCKET


def synthetic_resume(directory: Path) -> tuple[MasterResume, Path, Path, Path]:
    master = MasterResume.model_validate(
        {
            "contact": {"name": "Synthetic Candidate"},
            "summary": "Backend developer with Python & SQL experience.",
            "approved_summary_rephrasings": ["Backend developer using Python & SQL."],
            "skills": [
                {
                    "name": "Languages",
                    "items": ["Python", "SQL"],
                    "raw": "Languages: Python, SQL",
                }
            ],
            "experience": [
                {
                    "id": "role-1",
                    "employer": "Example Labs",
                    "title": "Backend Developer",
                    "location": "Lagos",
                    "dates": "2021–2023",
                    "bullets": [
                        {
                            "id": "b-1",
                            "text": "Built Python APIs serving 20 users.",
                            "approved_rephrasings": ["Built Python APIs for 20 users."],
                        },
                        {"id": "b-2", "text": "Used SQL to reduce response by 15%."},
                    ],
                }
            ],
            "projects": [
                {
                    "id": "p-1",
                    "name": "Example CLI",
                    "bullets": [{"id": "pb-1", "text": "Built a CLI with Python."}],
                },
                {
                    "id": "p-2",
                    "name": "Example Tool",
                    "bullets": [{"id": "pb-2", "text": "Used SQL for reports."}],
                },
            ],
            "education": [
                {
                    "institution": "Example University",
                    "degree": "BSc Computing",
                    "dates": "2017–2021",
                }
            ],
            "certifications": [{"name": "Example Certificate"}],
            "reviewed": True,
        }
    )
    original = directory / "master.docx"
    document = Document()
    document.sections[0].left_margin = Inches(0.8)
    document.sections[0].header.paragraphs[0].text = "Private synthetic fixture"
    document.sections[0].footer.paragraphs[0].text = "Fixture footer"
    document.add_paragraph(master.contact.name, "Title")
    document.add_heading("Summary", 1)
    paragraph = document.add_paragraph()
    first = paragraph.add_run("Backend developer ")
    first.bold = True
    first.font.name = "Arial"
    first.font.size = Pt(11)
    second = paragraph.add_run("with Python & SQL experience.")
    second.italic = True
    second.font.color.rgb = RGBColor(30, 50, 70)
    document.add_heading("Skills", 1)
    document.add_table(rows=1, cols=1).cell(0, 0).text = master.skills[0].raw
    document.add_heading("Experience", 1)
    document.add_paragraph("Example Labs | Backend Developer | Lagos | 2021–2023")
    for bullet in master.experience[0].bullets:
        document.add_paragraph(bullet.text, "List Bullet")
    document.add_heading("Projects", 1)
    for project in master.projects:
        document.add_paragraph(project.name, "Heading 2")
        for bullet in project.bullets:
            document.add_paragraph(bullet.text, "List Bullet")
    document.add_heading("Education", 1)
    document.add_paragraph("Example University | BSc Computing | 2017–2021")
    document.add_heading("Certifications", 1)
    document.add_paragraph("Example Certificate")
    document.save(original)
    template, layout = directory / "template.docx", directory / "layout.json"
    create_template(original, template, layout, master)
    return master, original, template, layout


def job() -> dict[str, Any]:
    return {
        "id": str(uuid4()),
        "company": "Example Company",
        "title": "Backend Engineer",
        "description_text": "Requirements: Python, SQL, and AWS.",
        "content_hash": "a" * 64,
    }


class ControlledBackend:
    def __init__(self, outputs: list[Any]) -> None:
        self.outputs = list(outputs)
        self.calls: list[dict[str, Any]] = []

    def complete(self, system: str, content: str) -> Reply:
        assert "untrusted" in system and "JSON" in system
        self.calls.append(json.loads(content))
        value = self.outputs.pop(0)
        if isinstance(value, Exception):
            raise value
        return Reply(value if isinstance(value, str) else json.dumps(value), 100, 30)


class PgResumeStore:
    def __init__(self, pg: PostgresStore, upload_error: bool = False) -> None:
        self.pg = pg
        self.uploads: dict[str, bytes] = {}
        self.upload_error = upload_error

    def assert_private(self) -> None:
        assert (
            self.pg.db.query(
                "select public from storage.buckets where id=$1", [BUCKET]
            )[0]["public"]
            is False
        )

    def claim(self, claim_id: str, include_maybe: bool) -> dict[str, Any] | None:
        return self.pg.db.query(
            "select public.radar_claim_tailor($1::uuid,$2::boolean) as result",
            [claim_id, include_maybe],
        )[0]["result"]

    def upload(self, path: str, document: Path) -> None:
        if self.upload_error:
            from pipeline.storage import StoreError

            raise StoreError("Private resume upload failed")
        assert path not in self.uploads
        self.uploads[path] = document.read_bytes()

    def remove(self, path: str) -> None:
        del self.uploads[path]

    def finish(
        self,
        job: dict[str, Any],
        claim_id: str,
        status: str,
        path: str | None,
        filename: str | None,
        notes: dict[str, Any] | None,
        error: str | None,
        metadata: dict[str, Any],
    ) -> bool:
        return self.pg.db.query(
            "select public.radar_finish_tailor($1::uuid,$2::uuid,$3::text,$4::text,"
            "$5::text,$6::text,$7::jsonb,$8::text,$9::jsonb) as result",
            [
                job["id"],
                claim_id,
                job["content_hash"],
                status,
                path,
                filename,
                json.dumps(notes) if notes is not None else None,
                error,
                json.dumps(metadata),
            ],
        )[0]["result"]
