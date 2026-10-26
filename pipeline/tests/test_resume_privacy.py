"""Synthetic contact sentinels must never cross the LLM HTTP boundary."""

import json
from copy import deepcopy
from pathlib import Path
from typing import Any

import httpx
import pytest
from docx import Document
from resume_helpers import (
    ControlledBackend,
    job,
    synthetic_resume,
    tailoring_proposal,
    verified_pagination,
)
from test_groq import groq_reply

from pipeline.llm import client as llm_module
from pipeline.llm.client import SDKBackend, Usage
from pipeline.resume.layout import LayoutError, compare, create_template, render
from pipeline.resume.privacy import ContactRedactor, tailoring_facts
from pipeline.resume.tailor import TailoringClient, tailor_job
from pipeline.resume.validate import unchanged
from pipeline.storage import PacedTransport

NAME = "Ada PrivacySentinel"
EMAIL = "ada-private@example.test"
PHONE = "+234 912 345 6789"
ADDRESS = "47 Secret Avenue, PrivateCity"
LINK = "https://portfolio.example.test/ada-private"
SENTINELS = [NAME, "Ada", "PrivacySentinel", EMAIL, PHONE, ADDRESS, LINK]


@pytest.mark.parametrize("provider", ["openai", "anthropic", "groq", "google"])
@pytest.mark.parametrize("raw_only", [False, True])
def test_no_contact_in_any_http_request_including_retries_and_local_render(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    recorded: dict[str, Any],
    make_http: Any,
    provider: str,
    raw_only: bool,
) -> None:
    master, original, template, layout = synthetic_resume(tmp_path)
    master.contact.name = NAME
    master.contact.lines = [f"{EMAIL} | {PHONE} | {LINK}", ADDRESS]
    if not raw_only:
        master.contact.email = EMAIL
        master.contact.phone = PHONE
        master.contact.location = ADDRESS
        master.contact.links = [LINK]
    # Personal data also occurs OUTSIDE the contact object, as it can in imports.
    private_line = "Project author: " + " | ".join(SENTINELS)
    master.projects[0].detail_lines = [private_line]
    document = Document(original)
    document.paragraphs[0].text = NAME
    for line in master.contact.lines:
        document.sections[0].header.add_paragraph(line)
    document.add_paragraph(private_line)
    document.save(original)
    create_template(original, template, layout, master)
    current_job = job()
    current_job["description_text"] += " Contact: " + private_line
    current_job["description_text"] += " Phone alternative: +234-912-345-6789."
    original_bullet = master.experience[0].bullets[0].text
    proposal = tailoring_proposal(
        master,
        bullet_rewrites=[
            {
                "id": master.experience[0].bullets[0].id,
                "text": original_bullet + " Improved by 99%.",
            }
        ],
    )
    outputs = [proposal, {"text": original_bullet}]
    calls: list[dict[str, Any]] = []

    def handler(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        calls.append(body)
        serialized = json.dumps(body, ensure_ascii=False).casefold()
        for sentinel in [*SENTINELS, "+234-912-345-6789"]:
            assert sentinel.casefold() not in serialized
        if provider == "google":
            sent_payload = json.loads(body["contents"][-1]["parts"][0]["text"])
        else:
            sent_payload = json.loads(body["messages"][-1]["content"])
        if "master" in sent_payload:
            assert "contact" not in sent_payload["master"]
        else:
            assert set(sent_payload) == {"slot", "job"}
        if len(calls) == 1:
            return httpx.Response(429, headers={"Retry-After": "1"})
        content = json.dumps(outputs.pop(0))
        if provider == "anthropic":
            return httpx.Response(
                200,
                json={
                    "id": "fixture-message",
                    "type": "message",
                    "role": "assistant",
                    "model": "fixture-model",
                    "content": [{"type": "text", "text": content}],
                    "stop_reason": "end_turn",
                    "usage": {"input_tokens": 10, "output_tokens": 5},
                },
            )
        if provider == "google":
            return httpx.Response(
                200,
                json={
                    "candidates": [
                        {
                            "content": {"parts": [{"text": content}]},
                            "finishReason": "STOP",
                        }
                    ],
                    "usageMetadata": {
                        "promptTokenCount": 10,
                        "candidatesTokenCount": 5,
                    },
                },
            )
        return httpx.Response(200, json=groq_reply(recorded, content))

    monkeypatch.setattr(
        llm_module,
        "PacedTransport",
        lambda http: PacedTransport(http, httpx.MockTransport(handler)),
    )
    http = make_http(lambda _: pytest.fail("No public HTTP expected"))
    client = TailoringClient(
        SDKBackend(provider, "synthetic-key", "fixture-model", http),
        sleep=lambda _: None,
    )
    result = tailor_job(
        master,
        current_job,
        original,
        template,
        layout,
        tmp_path / "output",
        client,
        Usage(),
        verified_pagination(),
    )
    assert result.status == "ready", result.error
    assert len(calls) == 3
    if provider == "google":
        first = json.loads(calls[0]["contents"][-1]["parts"][0]["text"])
        last = json.loads(calls[-1]["contents"][-1]["parts"][0]["text"])
    else:
        first = json.loads(calls[0]["messages"][-1]["content"])
        last = json.loads(calls[-1]["messages"][-1]["content"])
    assert last["slot"]["kind"] == "bullet"
    assert "validator_errors" not in last
    assert set(first["master"]) == {
        "summary",
        "approved_summary_rephrasings",
        "skills",
        "experience",
        "projects",
    }
    assert set(first["job"]) == {"title", "company", "description_text"}
    assert set(first["master"]["experience"][0]) == {"id", "bullets"}
    assert set(first["master"]["skills"][0]) == {"name", "items"}
    assert set(first) == {
        "master",
        "job",
        "validator_errors",
        "slot_character_limits",
    }
    assert first["slot_character_limits"]["bullets"]["b-1"] == len(original_bullet)
    assert "url" not in first["master"]["projects"][0]
    for immutable in (
        master.experience[0].employer,
        master.experience[0].dates,
        master.education[0].institution,
        master.education[0].degree,
        master.certifications[0].name,
    ):
        assert immutable not in json.dumps(first)
    assert "[private]" in first["master"]["projects"][0]["detail_lines"][0]
    assert compare(original, result.path, include_text=False) == []
    rendered = Document(result.path)
    text = "\n".join(p.text for p in rendered.paragraphs)
    text += "\n" + "\n".join(p.text for p in rendered.sections[0].header.paragraphs)
    for sentinel in SENTINELS:
        assert sentinel in text


def test_validator_feedback_is_redacted_before_the_next_request(tmp_path: Path) -> None:
    master, _, _, _ = synthetic_resume(tmp_path)
    master.contact.email = EMAIL
    valid = tailoring_proposal(master)
    backend = ControlledBackend([valid, valid])
    checks = 0

    def check(_: Any) -> None:
        nonlocal checks
        checks += 1
        if checks == 1:
            raise LayoutError("Synthetic validator detail: " + EMAIL)

    TailoringClient(backend, sleep=lambda _: None).propose(
        master, job(), Usage(), check
    )
    assert len(backend.calls) == 2
    assert EMAIL not in json.dumps(backend.calls)
    assert "[private]" in backend.calls[-1]["validator_errors"][0]


def test_contact_binding_uses_local_json_not_cached_template_text(
    tmp_path: Path,
) -> None:
    master, original, template, layout_path = synthetic_resume(tmp_path)
    layout = json.loads(layout_path.read_text())
    block = next(b for b in layout["blocks"] if b["binding"]["kind"] == "contact")
    assert block["binding"]["field"] == "name"
    # If cached contact text is corrupted, rendering still uses the local master.
    block["text"] = "x" * len(master.contact.name)
    for run in block["runs"]:
        run["text"] = "x" * len(run["text"])
    layout_path.write_text(json.dumps(layout))
    destination = tmp_path / "rendered.docx"
    render(master, unchanged(master), original, template, layout_path, destination)
    assert Document(destination).paragraphs[0].text == master.contact.name
    assert compare(original, destination) == []


def test_redaction_covers_contact_in_keys_and_escaped_unicode() -> None:
    from pipeline.resume.schema import Contact

    redactor = ContactRedactor(Contact(name="Chloé PrivacySentinel", email=EMAIL))
    payload = {EMAIL: "CHLOÉ\nPRIVACYSENTINEL " + EMAIL}
    redacted = redactor.payload(deepcopy(payload))
    assert redacted == {"[private]": "[private] [private]"}


def test_concatenated_personal_identifiers_are_redacted() -> None:
    from pipeline.resume.schema import Contact

    redactor = ContactRedactor(Contact(name=NAME, phone=PHONE))
    payload = {"roleAdaPrivacySentinel": "bullet+2349123456789"}
    assert redactor.payload(payload) == {"role[private][private]": "bullet[private]"}


def test_allowlisted_evidence_cannot_leak_education_through_colliding_ids(
    tmp_path: Path,
) -> None:
    master, _, _, _ = synthetic_resume(tmp_path)
    master.experience[0].bullets[0].id = "education:0:degree"
    _, evidence = tailoring_facts(master)
    assert evidence["education:0:degree"] == "/experience/0/bullets/0/text"
    assert master.education[0].degree not in evidence.values()
