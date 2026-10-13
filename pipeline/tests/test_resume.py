import json
from copy import deepcopy
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from pydantic import ValidationError
from resume_helpers import ControlledBackend, job, synthetic_resume

from pipeline.llm.client import Usage
from pipeline.resume.import_pdf import parse_pdf_text, split_skills
from pipeline.resume.init import main as init_main
from pipeline.resume.layout import (
    LayoutError,
    compare,
    create_template,
    draft_master,
    render,
)
from pipeline.resume.pdf import Pagination
from pipeline.resume.schema import SkillGroup, Tailoring
from pipeline.resume.tailor import TailoringClient, filename, tailor_job
from pipeline.resume.validate import TruthError, unchanged, validate_tailoring


@pytest.mark.parametrize(
    "mutation",
    [
        "skill",
        "metric",
        "currency",
        "bullet_id",
        "wrong_role",
        "summary_overflow",
        "bullet_overflow",
        "bullet_count",
        "technology",
        "proper_noun",
        "seniority",
        "unapproved",
        "control",
    ],
)
def test_validator_rejects_unsupported_facts_and_overflow(
    tmp_path: Path, mutation: str
) -> None:
    master, _, _, _ = synthetic_resume(tmp_path)
    payload = unchanged(master).model_dump()
    if mutation == "skill":
        payload["skills_order"].append("AWS")
    elif mutation in {"metric", "currency"}:
        payload["experience"][0]["bullets"][0]["text"] = (
            "Built APIs serving 200 users."
            if mutation == "metric"
            else "Built APIs earning $20."
        )
    elif mutation in {"bullet_id", "wrong_role"}:
        payload["experience"][0]["bullets"][0]["id"] = (
            "unknown" if mutation == "bullet_id" else "pb-1"
        )
    elif mutation == "summary_overflow":
        payload["summary"] = "x" * 1000
    elif mutation == "bullet_overflow":
        payload["experience"][0]["bullets"][0]["text"] = "x" * 1000
    elif mutation == "bullet_count":
        payload["experience"][0]["bullets"].append(
            payload["experience"][0]["bullets"][0]
        )
    elif mutation == "control":
        payload["summary"] = "\nInjected\fPage"
    else:
        text = {
            "technology": "Built APIs using AWS.",
            "proper_noun": "Built APIs for InventedCorp.",
            "seniority": "Led experts building APIs.",
            "unapproved": "Never used Python.",
        }[mutation]
        payload["experience"][0]["bullets"][0]["text"] = text
    with pytest.raises(TruthError):
        validate_tailoring(
            master, Tailoring.model_validate(payload), job()["description_text"]
        )


@pytest.mark.parametrize(
    "field", ["dates", "employer", "title", "education", "certifications"]
)
def test_llm_cannot_edit_immutable_fields(tmp_path: Path, field: str) -> None:
    master, _, _, _ = synthetic_resume(tmp_path)
    payload = unchanged(master).model_dump()
    target = (
        payload["experience"][0] if field in {"dates", "employer", "title"} else payload
    )
    target[field] = "Edited fact"
    with pytest.raises(ValidationError):
        Tailoring.model_validate_json(json.dumps(payload))


def test_approved_wording_still_cannot_add_metrics_or_technologies(
    tmp_path: Path,
) -> None:
    master, _, _, _ = synthetic_resume(tmp_path)
    for text in (
        "Built Python APIs for 200 users.",
        "Built AWS APIs for 20 users.",
        "Built APIs for InventedCorp.",
    ):
        modified = master.model_copy(deep=True)
        modified.experience[0].bullets[0].approved_rephrasings.append(text)
        payload = unchanged(modified).model_dump()
        payload["experience"][0]["bullets"][0]["text"] = text
        with pytest.raises(TruthError):
            validate_tailoring(
                modified, Tailoring.model_validate(payload), job()["description_text"]
            )


def test_valid_case_insensitive_skills_and_evidence_based_gaps(tmp_path: Path) -> None:
    master, _, _, _ = synthetic_resume(tmp_path)
    payload = unchanged(master).model_dump()
    payload["skills_order"] = ["sql", "PYTHON"]
    payload["match_notes"] = {
        "matched": [{"requirement": "Python", "evidence_ids": ["b-1"]}],
        "gaps": ["AWS"],
    }
    accepted = validate_tailoring(
        master, Tailoring.model_validate(payload), job()["description_text"]
    )
    assert accepted.skills_order == ["SQL", "Python"]
    assert accepted.match_notes.gaps == ["AWS"]
    for notes in (
        {"matched": [{"requirement": "AWS", "evidence_ids": ["b-1"]}], "gaps": []},
        {"matched": [], "gaps": ["Python"]},
        {"matched": [], "gaps": ["Invented requirement"]},
    ):
        payload["match_notes"] = notes
        with pytest.raises(TruthError):
            validate_tailoring(
                master, Tailoring.model_validate(payload), job()["description_text"]
            )


def test_unmodified_master_round_trip_preserves_paragraphs_runs_and_styles(
    tmp_path: Path,
) -> None:
    master, original, template, layout = synthetic_resume(tmp_path)
    rendered = tmp_path / "roundtrip.docx"
    render(master, unchanged(master), original, template, layout, rendered)
    assert compare(original, rendered) == []
    before, after = Document(original), Document(rendered)
    assert [(p.text, p.style.style_id) for p in before.paragraphs] == [
        (p.text, p.style.style_id) for p in after.paragraphs
    ]
    assert before.tables[0].cell(0, 0).text == after.tables[0].cell(0, 0).text
    assert (
        before.sections[0].header.paragraphs[0].text
        == after.sections[0].header.paragraphs[0].text
    )


def test_valid_approved_output_renders_and_immutable_facts_stay_copied(
    tmp_path: Path,
) -> None:
    master, original, template, layout = synthetic_resume(tmp_path)
    payload = unchanged(master).model_dump()
    payload["summary"] = master.approved_summary_rephrasings[0]
    payload["experience"][0]["bullets"][0]["text"] = (
        master.experience[0].bullets[0].approved_rephrasings[0]
    )
    accepted = validate_tailoring(
        master, Tailoring.model_validate(payload), job()["description_text"]
    )
    rendered = tmp_path / "valid.docx"
    render(master, accepted, original, template, layout, rendered)
    assert compare(original, rendered, include_text=False) == []
    text = "\n".join(p.text for p in Document(rendered).paragraphs)
    assert (
        accepted.summary in text
        and "Example Labs | Backend Developer | Lagos | 2021–2023" in text
    )
    assert "BSc Computing" in text and "Example Certificate" in text


def test_projects_reorder_as_whole_blocks_and_selected_bullets_remove_cleanly(
    tmp_path: Path,
) -> None:
    master, original, template, layout = synthetic_resume(tmp_path)
    payload = unchanged(master).model_dump()
    payload["projects_order"].reverse()
    rendered = tmp_path / "reordered.docx"
    render(
        master,
        validate_tailoring(master, Tailoring.model_validate(payload), ""),
        original,
        template,
        layout,
        rendered,
    )
    text = [p.text for p in Document(rendered).paragraphs]
    assert text.index("Example Tool") < text.index("Example CLI")
    payload = unchanged(master).model_dump()
    payload["experience"][0]["bullet_order"] = ["b-2"]
    payload["experience"][0]["bullets"] = [payload["experience"][0]["bullets"][1]]
    render(
        master,
        validate_tailoring(master, Tailoring.model_validate(payload), ""),
        original,
        template,
        layout,
        rendered,
    )
    text = [p.text for p in Document(rendered).paragraphs]
    assert master.experience[0].bullets[0].text not in text
    assert master.experience[0].bullets[1].text in text


def test_retry_feedback_and_cost_count_every_attempt(tmp_path: Path) -> None:
    master, original, template, layout = synthetic_resume(tmp_path)
    good = unchanged(master).model_dump()
    bad = deepcopy(good)
    bad["skills_order"].append("AWS")
    backend = ControlledBackend([bad, good])
    client = TailoringClient(backend, 1, 2, sleep=lambda _: None)
    usage = Usage()
    result = tailor_job(
        master,
        job(),
        original,
        template,
        layout,
        tmp_path / "output",
        client,
        usage,
        Pagination(discover=False),
    )
    assert result.status == "ready" and usage.attempts == 2
    assert usage.input_tokens == 200 and usage.output_tokens == 60
    assert usage.cost_usd == pytest.approx(0.00032)
    assert "invented skill" in backend.calls[1]["validator_errors"][0]
    assert "contact" not in backend.calls[0]["master"]


def test_validation_failure_falls_back_to_byte_identical_master(tmp_path: Path) -> None:
    master, original, template, layout = synthetic_resume(tmp_path)
    bad = unchanged(master).model_dump()
    bad["skills_order"].append("AWS")
    backend = ControlledBackend([bad] * 3)
    usage = Usage()
    result = tailor_job(
        master,
        job(),
        original,
        template,
        layout,
        tmp_path / "output",
        TailoringClient(backend, sleep=lambda _: None),
        usage,
        Pagination(discover=False),
    )
    assert result.status == "failed" and "3 attempts" in result.error
    assert result.path.read_bytes() == original.read_bytes()
    assert usage.attempts == 3 and result.match_notes is None


def test_no_libreoffice_rejects_unverified_changed_layout(tmp_path: Path) -> None:
    master, original, template, layout = synthetic_resume(tmp_path)
    proposed = unchanged(master).model_dump()
    proposed["summary"] = master.approved_summary_rephrasings[0]
    result = tailor_job(
        master,
        job(),
        original,
        template,
        layout,
        tmp_path / "output",
        TailoringClient(ControlledBackend([proposed] * 3), sleep=lambda _: None),
        Usage(),
        Pagination(discover=False),
    )
    assert result.status == "failed" and "page count" in result.error
    assert result.path.read_bytes() == original.read_bytes()


def test_pdf_page_count_guard_rejects_overflow(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    master, original, template, layout = synthetic_resume(tmp_path)
    rendered = tmp_path / "candidate.docx"
    render(master, unchanged(master), original, template, layout, rendered)
    engine = Pagination("controlled-test-executable", discover=False)
    monkeypatch.setattr(
        engine, "convert", lambda source, dest: 1 if source == original else 2
    )
    with pytest.raises(LayoutError, match="page count"):
        engine.check(original, rendered)


def test_offline_init_and_render_check_need_no_llm(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: Any
) -> None:
    master, original, _, _ = synthetic_resume(tmp_path)
    for name in ("LLM_PROVIDER", "LLM_MODEL", "LLM_API_KEY"):
        monkeypatch.delenv(name, raising=False)
    args = [
        "--master-docx",
        str(original),
        "--master-json",
        str(tmp_path / "imported.json"),
        "--template",
        str(tmp_path / "imported.docx"),
        "--layout",
        str(tmp_path / "imported-layout.json"),
        "--output",
        str(tmp_path / "checked.docx"),
    ]
    assert init_main(args) == 0
    assert "Round-trip" in capsys.readouterr().out
    extracted = draft_master(original)
    assert not extracted.reviewed and len(extracted.experience[0].bullets) == 2
    assert len(extracted.projects) == 2
    assert init_main([*args, "--render-check"]) == 0


def test_pdf_wrapping_and_combined_skill_groups_preserve_literal_baseline(
    tmp_path: Path,
) -> None:
    master, original, template, layout = synthetic_resume(tmp_path)
    document = Document(original)
    summary = next(p for p in document.paragraphs if p.text == master.summary)
    summary.text = "Backend developer "
    continuation = document.add_paragraph("with Python & SQL experience.")
    summary._p.addnext(continuation._p)
    master.skills.append(SkillGroup(name="Tools", items=["Git"], raw="Tools: Git"))
    document.tables[0].cell(0, 0).text += "   Tools: Git"
    document.save(original)
    create_template(original, template, layout, master)
    output = tmp_path / "roundtrip.docx"
    render(master, unchanged(master), original, template, layout, output)
    assert compare(original, output) == []
    proposed = unchanged(master).model_copy(
        update={
            "summary": "Backend developer using Python & SQL.",
            "skills_order": ["SQL", "Python", "Git"],
        }
    )
    render(master, proposed, original, template, layout, output)
    text = " ".join(p.text for p in Document(output).paragraphs)
    assert proposed.summary in text
    assert Document(output).tables[0].cell(0, 0).text == (
        "Languages: SQL, Python   Tools: Git"
    )


def test_master_mutation_and_template_tampering_fail_closed(tmp_path: Path) -> None:
    master, original, template, layout = synthetic_resume(tmp_path)
    changed = master.model_copy(deep=True)
    changed.experience[0].dates = "2099–2100"
    with pytest.raises(LayoutError):
        render(
            changed,
            unchanged(changed),
            original,
            template,
            layout,
            tmp_path / "bad.docx",
        )
    with pytest.raises(LayoutError, match="facts disagree"):
        create_template(original, tmp_path / "new.docx", tmp_path / "new.json", changed)
    template.write_bytes(b"tampered")
    with pytest.raises(LayoutError):
        render(
            master, unchanged(master), original, template, layout, tmp_path / "bad.docx"
        )


def test_pdf_draft_groups_columns_and_preserves_ambiguous_dates() -> None:
    source = """Synthetic Candidate
SUMMARY
Backend developer with Python.
EXPERIENCE
Backend Developer — Example Labs (Remote)        Jun 2026 – PresentJuly 2026
● Built APIs with Python
  serving 20 users.
PROJECTS
Example CLI        Jan 2026 – Present
Python, SQL   |   Source: example.invalid
● Built a tool.
TECHNICAL SKILLS
Languages: Python, SQL (PostgreSQL, MySQL)
EDUCATION
Example University — Lagos        2021
BSc Computing
"""
    master, issues = parse_pdf_text(source)
    assert not master.reviewed and len(master.projects) == 1
    assert master.experience[0].dates == "Jun 2026 – PresentJuly 2026"
    assert (
        master.experience[0].bullets[0].text
        == "Built APIs with Python serving 20 users."
    )
    assert any("overlapping" in issue for issue in issues)
    assert split_skills("Python, SQL (PostgreSQL, MySQL)") == [
        "Python",
        "SQL (PostgreSQL, MySQL)",
    ]


def test_filename_is_sanitized_and_bounded(tmp_path: Path) -> None:
    master, _, _, _ = synthetic_resume(tmp_path)
    name = filename(master, "../Company:Bad", "Role/Title" * 100)
    assert "/" not in name and ":" not in name and ".." not in name
    assert name.endswith(".docx") and len(name) <= 125


def test_education_and_project_facts_are_available_as_note_evidence(
    tmp_path: Path,
) -> None:
    master, _, _, _ = synthetic_resume(tmp_path)
    payload = unchanged(master).model_dump()
    description = "Requirements: BSc Computing. Prior work with an Example CLI."
    payload["match_notes"] = {
        "matched": [
            {"requirement": "BSc Computing", "evidence_ids": ["education:0:degree"]},
            {"requirement": "Example CLI", "evidence_ids": ["project:p-1:name"]},
        ],
        "gaps": [],
    }
    validate_tailoring(master, Tailoring.model_validate(payload), description)
    payload["match_notes"] = {"matched": [], "gaps": ["BSc Computing"]}
    with pytest.raises(TruthError, match="gap conflicts"):
        validate_tailoring(master, Tailoring.model_validate(payload), description)
