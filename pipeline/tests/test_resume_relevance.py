from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from docx import Document
from resume_helpers import ControlledBackend, synthetic_resume, tailoring_proposal

from pipeline.llm.client import Usage
from pipeline.resume import tailor as tailor_module
from pipeline.resume.layout import LayoutError
from pipeline.resume.layout import render as render_resume
from pipeline.resume.pdf import Pagination
from pipeline.resume.relevance import deterministic_tailoring, matching_skills
from pipeline.resume.schema import TailoringProposal
from pipeline.resume.tailor import TailoringClient, tailor_job
from pipeline.resume.validate import audit_proposal
from pipeline.run_tailor import Outcome, print_outcome


def tailored(
    tmp_path: Path,
    *,
    skills: list[dict[str, Any]],
    bullets: list[str],
    title: str,
    description: str,
):
    master, original, template, layout = synthetic_resume(
        tmp_path, skills=skills, experience_bullets=bullets
    )
    pagination = Pagination("controlled-renderer", discover=False)
    pagination.convert = lambda *_: 1
    client = TailoringClient(
        ControlledBackend([tailoring_proposal(master)]), sleep=lambda _: None
    )
    result = tailor_job(
        master,
        {
            "id": "004ebdc0-ca87-42a4-924f-7a2f7f5c2982",
            "company": "Example Company",
            "title": title,
            "description_text": description,
        },
        original,
        template,
        layout,
        tmp_path / "output",
        client,
        Usage(),
        pagination,
    )
    return master, result, client


def test_backend_job_gets_role_summary_skill_groups_and_bullets_reordered(
    tmp_path: Path, capsys: Any
) -> None:
    master, result, client = tailored(
        tmp_path,
        skills=[
            {"name": "Languages", "items": ["Python", "JavaScript/TypeScript"]},
            {"name": "Backend", "items": ["FastAPI", "RESTful APIs"]},
            {"name": "Databases", "items": ["PostgreSQL", "Redis"]},
        ],
        bullets=[
            "Designed PostgreSQL indexes to improve query speed.",
            "Built FastAPI backend services in Python.",
            "Used Redis to cache application responses.",
        ],
        title="FastAPI Backend Engineer",
        description="Required: FastAPI, Python, PostgreSQL. Build RESTful APIs.",
    )
    assert result.status == "ready"
    assert client.last_diagnostic is None
    audit = audit_proposal(
        master,
        TailoringProposal.model_validate(tailoring_proposal(master)).model_dump_json(),
        "Required: FastAPI, Python, PostgreSQL. Build RESTful APIs.",
    )
    assert audit["schema"]["accepted"] is True
    assert audit["fields"]["summary"]["changed"] is False
    assert result.changes["tailoring_changed"] is True
    assert result.change_report["summary_before"] == master.summary
    assert "FastAPI Backend Engineer" in result.change_report["summary_after"]
    assert result.change_report["skills_after"][0]["group"] == "Backend"
    assert result.change_report["matched_requirements"]
    document = Document(result.path)
    assert document.tables[0].cell(0, 0).text.startswith("Backend: FastAPI")
    bullet_text = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
    assert bullet_text[0] == "Built FastAPI backend services in Python.", result.changes

    print_outcome(
        Outcome(
            "job-1",
            result.status,
            Usage(),
            path=result.path,
            report_path=result.report_path,
            changes=result.changes,
            warnings=result.warnings,
            change_report=result.change_report,
        )
    )
    printed = capsys.readouterr().out
    assert "Summary before:" in printed and "Summary after:" in printed
    assert "Skills before:" in printed and "Skills after:" in printed
    assert "Matched requirements:" in printed and "Gaps:" in printed


def test_node_react_job_reorders_frontend_skills_and_relevant_bullets(
    tmp_path: Path,
) -> None:
    _, result, _ = tailored(
        tmp_path,
        skills=[
            {"name": "Languages", "items": ["Python", "JavaScript/TypeScript"]},
            {"name": "Frontend", "items": ["React", "Next.js"]},
            {"name": "Backend", "items": ["FastAPI", "RESTful APIs"]},
        ],
        bullets=[
            "Built Python APIs for a small dashboard system.",
            "Developed React screens with TypeScript.",
            "Used SQL query indexes to improve report speed.",
        ],
        title="Full Stack Engineer (Node.js & React)",
        description=(
            "Required skills: React, TypeScript, and SQL. "
            "Build accessible React components."
        ),
    )
    assert result.status == "ready"
    assert result.changes["tailoring_changed"] is True
    assert result.change_report["skills_after"][0]["group"] == "Frontend", (
        result.change_report["skills_after"]
    )
    assert result.change_report["skills_after"][0]["skills"][0] == "React"
    assert result.change_report["summary_after"].startswith(
        "Targeting Full Stack Engineer"
    )
    document = Document(result.path)
    assert document.tables[0].cell(0, 0).text.startswith("Frontend: React")
    bullet_text = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
    assert bullet_text[0] == "Developed React screens with TypeScript.", result.changes
    project_names = [p.text for p in document.paragraphs if p.style.name == "Heading 2"]
    assert project_names.index("Example Tool") < project_names.index("Example CLI")


def test_aliases_match_node_and_postgres_spellings_to_master_skills(
    tmp_path: Path,
) -> None:
    master, _, _, _ = synthetic_resume(
        tmp_path,
        skills=[
            {"name": "Backend", "items": ["Node.js", "PostgreSQL"]},
            {"name": "Frontend", "items": ["React"]},
        ],
    )
    job = {
        "title": "Junior Node Developer",
        "description_text": "Use Postgres and React to build APIs.",
    }
    assert matching_skills(master, job) == ["Node.js", "PostgreSQL", "React"]
    ordered = deterministic_tailoring(master, job)
    assert ordered.skills_order[:2] == ["Node.js", "PostgreSQL"]


def test_overlong_summary_retry_reverts_only_summary_and_keeps_ordering(
    tmp_path: Path,
) -> None:
    master, original, template, layout = synthetic_resume(
        tmp_path,
        skills=[
            {"name": "Languages", "items": ["Python"]},
            {"name": "Backend", "items": ["FastAPI", "PostgreSQL"]},
        ],
        experience_bullets=[
            "Designed PostgreSQL indexes to improve query speed.",
            "Built FastAPI backend services in Python.",
        ],
    )
    proposal = tailoring_proposal(master, summary="Overlong summary proposal " * 5)
    backend = ControlledBackend(
        [proposal, {"text": "Still too long to fit this summary slot " * 3}]
    )
    pagination = Pagination("controlled-renderer", discover=False)
    pagination.convert = lambda *_: 1
    result = tailor_job(
        master,
        {
            "id": "004ebdc0-ca87-42a4-924f-7a2f7f5c2982",
            "company": "Example Company",
            "title": "Backend Engineer",
            "description_text": "Required: PostgreSQL and FastAPI.",
        },
        original,
        template,
        layout,
        tmp_path / "output",
        TailoringClient(backend, sleep=lambda _: None),
        Usage(),
        pagination,
    )
    assert result.status == "ready"
    assert result.change_report["summary_after"] == master.summary, result.warnings
    assert result.change_report["skills_after"][0]["group"] == "Backend"
    assert "kept the original summary" in " ".join(result.warnings)
    document = Document(result.path)
    bullets = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
    assert bullets[0] == "Built FastAPI backend services in Python."
    assert len(backend.calls) == 2


def test_local_layout_overflow_keeps_other_safe_tailoring_changes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    master, original, template, layout = synthetic_resume(
        tmp_path,
        skills=[
            {"name": "Database Technologies", "items": ["SQL"]},
            {"name": "Backend", "items": ["Python", "FastAPI"]},
        ],
        experience_bullets=[
            "Used SQL database queries to optimize response handling by 15%.",
            "Built FastAPI backend services.",
        ],
    )
    original_order = [bullet.id for bullet in master.experience[0].bullets]

    def reject_reordered_role(*args: Any, **kwargs: Any) -> None:
        edit = args[1]
        if edit.experience[0].bullet_order != original_order:
            raise LayoutError("Synthetic bullet slot overflow")
        render_resume(*args, **kwargs)

    monkeypatch.setattr(tailor_module, "render", reject_reordered_role)
    pagination = Pagination("controlled-renderer", discover=False)
    pagination.convert = lambda *_: 1
    result = tailor_job(
        master,
        {
            "id": "004ebdc0-87c9-4d76-98a5-857628a65ac4",
            "company": "Example Company",
            "title": "FastAPI Backend Engineer",
            "description_text": "Required: FastAPI, Python, and SQL.",
        },
        original,
        template,
        layout,
        tmp_path / "output",
        None,
        Usage(),
        pagination,
        local_only=True,
    )

    assert result.status == "ready", result.error
    assert result.change_report["skills_after"][0]["group"] == "Backend"
    assert any("master experience role" in warning for warning in result.warnings)
    document = Document(result.path)
    assert document.tables[0].cell(0, 0).text.startswith("Backend: FastAPI")
    bullets = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
    assert bullets[:2] == [bullet.text for bullet in master.experience[0].bullets]


def test_overlong_bullet_retry_reverts_only_that_bullet(tmp_path: Path) -> None:
    master, original, template, layout = synthetic_resume(
        tmp_path,
        skills=[
            {"name": "Databases", "items": ["SQL"]},
            {"name": "Backend", "items": ["Python", "FastAPI"]},
        ],
        experience_bullets=[
            "Used SQL to reduce response by 15%.",
            "Built Python APIs serving 20 users.",
        ],
    )
    proposal = tailoring_proposal(
        master,
        bullet_rewrites=[
            {
                "id": "b-1",
                "text": (
                    "Used SQL to reduce response by 15% with a newly invented claim."
                ),
            }
        ],
    )
    backend = ControlledBackend(
        [proposal, {"text": "Still too long to fit the original bullet slot."}]
    )
    pagination = Pagination("controlled-renderer", discover=False)
    pagination.convert = lambda *_: 1
    result = tailor_job(
        master,
        {
            "id": "004ebdc0-ca87-42a4-924f-7a2f7f5c2982",
            "company": "Example Company",
            "title": "Python Backend Engineer",
            "description_text": "Required: Python and FastAPI.",
        },
        original,
        template,
        layout,
        tmp_path / "output",
        TailoringClient(backend, sleep=lambda _: None),
        Usage(),
        pagination,
    )
    assert result.status == "ready"
    assert any("b-1 exceeded its slot twice" in w for w in result.warnings)
    assert result.changes["bullet_roles_reordered"] == 1
    document = Document(result.path)
    bullets = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
    assert bullets[:2] == [
        "Built Python APIs serving 20 users.",
        "Used SQL to reduce response by 15%.",
    ]


def test_ios_without_matching_master_skills_uses_low_fit_master_without_llm(
    tmp_path: Path,
) -> None:
    master, original, template, layout = synthetic_resume(tmp_path)
    backend = ControlledBackend([])
    pagination = Pagination(discover=False)
    result = tailor_job(
        master,
        {
            "id": "02bef5df-4bda-482a-8b4c-f312c4d0361f",
            "company": "Example Company",
            "title": "iOS Developer",
            "description_text": "Requirements: Swift, UIKit, and iOS development.",
        },
        original,
        template,
        layout,
        tmp_path / "output",
        TailoringClient(backend),
        Usage(),
        pagination,
    )
    assert result.status == "low_fit"
    assert result.path.read_bytes() == original.read_bytes()
    assert backend.calls == []
    assert result.changes["tailoring_changed"] is False
    assert "No job skills meaningfully overlap" in result.warnings[0]
